"""AI generation, web-search, and uploaded-file extraction helpers."""

import io
import json
import sys
from pathlib import Path

import httpx
import openpyxl

# 确保项目根目录在 sys.path 中
_PROJECT_ROOT = str(Path(__file__).parent.parent)
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from config import settings  # noqa: E402


def search_company(project: str, category: str = '') -> str:
    """用 Tavily 检索项目公开信息，返回拼合后的文本摘要。无 key 或失败时静默返回空串。"""
    key = settings.TAVILY_API_KEY
    if not key or not project:
        return ''
    queries = [
        f'{project} 融资 投资 业务',
        f'{project} {category}'.strip() if category else f'{project} 产品 市场',
    ]
    seen, results = set(), []
    for q in queries:
        try:
            resp = httpx.post(
                f'{settings.TAVILY_BASE_URL}/search',
                json={'api_key': key, 'query': q,
                      'search_depth': 'basic', 'max_results': 4},
                timeout=15,
            )
            for item in resp.json().get('results', []):
                snippet = f"[{item.get('title','')}]\n{item.get('content','')[:400]}"
                if snippet not in seen:
                    seen.add(snippet)
                    results.append(snippet)
        except Exception:
            pass
    return '\n\n'.join(results[:6])


def extract_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in ('.txt', '.md', '.csv'):
        return content.decode('utf-8', errors='replace')
    if ext == '.pdf':
        for mod in ('pypdf', 'PyPDF2'):
            try:
                m = __import__(mod)
                reader = m.PdfReader(io.BytesIO(content))
                return '\n'.join(p.extract_text() or '' for p in reader.pages)
            except ImportError:
                continue
        return f'[PDF 提取失败，请安装 pip install pypdf | 文件：{filename}]'
    if ext == '.docx':
        try:
            from docx import Document
            return '\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)
        except ImportError:
            return f'[DOCX 提取失败，请安装 pip install python-docx | 文件：{filename}]'
    if ext in ('.xlsx', '.xls'):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f'[Sheet: {ws.title}]')
                for row in ws.iter_rows(values_only=True):
                    s = '\t'.join(str(v) if v is not None else '' for v in row)
                    if s.strip():
                        lines.append(s)
            wb.close()
            return '\n'.join(lines)
        except Exception as e:
            return f'[Excel 提取失败：{e}]'
    return content.decode('utf-8', errors='replace')


def stream_weekly_generation(project: str, week: str, extra: str, category: str,
                             biz_scope: str, stage: str, progress: str, file_data,
                             system_prompt: str, db_session=None):
    """Yield SSE-ready payload strings for the weekly-report generation endpoint."""
    from reporter.services.ai_provider import get_provider, get_ai_model

    try:
        provider = get_provider(db_session)
    except ValueError as e:
        yield f'data: {json.dumps({"error": str(e)})}\n\n'
        return

    # 获取当前配置的 model
    ai_model = get_ai_model(db_session)

    file_blocks = []
    for filename, content in file_data:
        try:
            text = extract_text(filename, content)
            if text.strip():
                file_blocks.append(f'=== {filename} ===\n{text[:10000]}')
        except Exception as e:
            yield f'data: {json.dumps({"error": f"{filename} 解析失败：{e}"})}\n\n'
            return

    tavily_key = settings.TAVILY_API_KEY
    if tavily_key:
        yield f'data: {json.dumps({"status": "正在检索公开信息…"})}\n\n'
    search_text = search_company(project, category)

    parts = [f'项目名称：{project}', f'周次：{week}']
    if category:
        parts.append(f'项目类别（细分）：{category}')
    if biz_scope:
        parts.append(f'行业大类：{biz_scope}')
    if stage:
        parts.append(f'当前阶段：{stage}')
    if progress:
        parts.append(f'\n--- 本周进展笔记（手动记录）---\n{progress}')
    if extra:
        parts.append(f'补充说明：{extra}')
    if search_text:
        parts.append(f'\n--- 公开信息（Tavily 检索）---\n{search_text}')
    if file_blocks:
        parts.append('\n--- 上传材料 ---\n' + '\n\n'.join(file_blocks))
    user_msg = '\n'.join(parts)

    for chunk in provider.stream_chat(system_prompt, user_msg, ai_model):
        yield chunk
